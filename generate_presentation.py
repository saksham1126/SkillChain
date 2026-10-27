import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Initialize 16:9 Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Theme Colors (Modern Professional Navy/Blue)
C_NAVY_DARK = RGBColor(15, 23, 42)      # #0F172A
C_NAVY_LIGHT = RGBColor(30, 41, 59)     # #1E293B
C_BLUE = RGBColor(37, 99, 235)          # #2563EB
C_BLUE_LIGHT = RGBColor(239, 246, 255)  # #EFF6FF
C_SLATE_BG = RGBColor(248, 250, 252)    # #F8FAFC
C_CARD_BG = RGBColor(255, 255, 255)     # #FFFFFF
C_BORDER = RGBColor(203, 213, 225)      # #CBD5E1
C_TEXT_MAIN = RGBColor(15, 23, 42)      # #0F172A
C_TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B
C_GREEN = RGBColor(5, 150, 105)         # #059669
C_WHITE = RGBColor(255, 255, 255)


def add_text_line(tf, text, font_size=11, bold=False, color=C_TEXT_MAIN, space_after=4, font_name="Calibri"):
    """
    Safely adds lines of text without ever embedding literal newlines inside
    an OpenXML <a:t> node, preventing PowerPoint corruption errors.
    """
    lines = text.split("\n")
    for idx, line in enumerate(lines):
        if len(tf.paragraphs) == 1 and tf.paragraphs[0].text == "":
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.name = font_name
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.color.rgb = color
        if idx == len(lines) - 1:
            p.space_after = Pt(space_after)
        else:
            p.space_after = Pt(2)


def add_header(slide, title_text, category="SKILLCHAIN • BLOCKCHAIN PROJECT"):
    """Draws a clean, safe header without invalid line XML."""
    # Top banner background
    banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.15))
    banner.fill.solid()
    banner.fill.fore_color.rgb = C_NAVY_DARK
    banner.line.color.rgb = C_NAVY_DARK
    banner.line.width = Pt(0)

    # Category Pill
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.18), Inches(10), Inches(0.3))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    add_text_line(tf_cat, category.upper(), font_size=9.5, bold=True, color=RGBColor(96, 165, 250))

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11), Inches(0.6))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    add_text_line(tf_title, title_text, font_size=20, bold=True, color=C_WHITE)

    # Bottom blue line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.13), Inches(13.333), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = C_BLUE
    line.line.color.rgb = C_BLUE
    line.line.width = Pt(0)


def add_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_BORDER):
    """Draws a clean modern rounded card with valid OpenXML borders."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.color.rgb = bg_color
        card.line.width = Pt(0)
    return card


# ==============================================================================
# SLIDE 1: Title Slide (Dark Theme)
# ==============================================================================
slide1 = prs.slides.add_slide(blank_layout)
bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = C_NAVY_DARK
bg1.line.color.rgb = C_NAVY_DARK
bg1.line.width = Pt(0)

dec1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(0.12), Inches(4.2))
dec1.fill.solid()
dec1.fill.fore_color.rgb = C_BLUE
dec1.line.color.rgb = C_BLUE
dec1.line.width = Pt(0)

tb1 = slide1.shapes.add_textbox(Inches(1.2), Inches(1.6), Inches(11), Inches(4.2))
tf1 = tb1.text_frame
tf1.word_wrap = True

add_text_line(tf1, "COLLEGE MAJOR PROJECT PRESENTATION • BLOCKCHAIN & WEB3", font_size=11, bold=True, color=RGBColor(96, 165, 250), space_after=14)
add_text_line(tf1, "SkillChain", font_size=46, bold=True, color=C_WHITE, space_after=8)
add_text_line(tf1, "Blockchain-Based Skill Passport & Decentralized Credential Verification", font_size=20, bold=True, color=RGBColor(226, 232, 240), space_after=16)
add_text_line(tf1, "A zero-trust academic verification platform leveraging Ethereum smart contracts and cryptographic SHA-256 hashing to eliminate certificate fraud.", font_size=13, color=RGBColor(148, 163, 184), space_after=28)
add_text_line(tf1, "Author: Saksham Pathak   |   Tech: Solidity, Hardhat, Ethers.js, React", font_size=12, bold=True, color=RGBColor(203, 213, 225))


# ==============================================================================
# SLIDE 2: Background & Motivation
# ==============================================================================
slide2 = prs.slides.add_slide(blank_layout)
add_header(slide2, "Background & Motivation")

col_w = Inches(3.7)
card_y = Inches(1.6)
card_h = Inches(5.2)

# Card 1
add_card(slide2, Inches(0.8), card_y, col_w, card_h)
tb = slide2.shapes.add_textbox(Inches(1.05), card_y + Inches(0.25), col_w - Inches(0.5), card_h - Inches(0.5))
tf = tb.text_frame
tf.word_wrap = True
add_text_line(tf, "The Credential Crisis", font_size=16, bold=True, color=C_NAVY_DARK, space_after=12)
points1 = [
    "Degrees, diplomas, and certifications are foundational for employment and admissions.",
    "Traditional paper documents and static digital PDFs are exceptionally easy to forge.",
    "Over 30% of resumes contain modified or embellished educational and skill claims.",
    "Credential fraud has become a multi-billion dollar worldwide challenge."
]
for pt in points1:
    add_text_line(tf, "- " + pt, font_size=11.5, color=C_TEXT_MAIN, space_after=10)

# Card 2
add_card(slide2, Inches(4.8), card_y, col_w, card_h)
tb = slide2.shapes.add_textbox(Inches(5.05), card_y + Inches(0.25), col_w - Inches(0.5), card_h - Inches(0.5))
tf = tb.text_frame
tf.word_wrap = True
add_text_line(tf, "Centralized Flaws", font_size=16, bold=True, color=C_NAVY_DARK, space_after=12)
points2 = [
    "Verification currently requires contacting registrar offices or proprietary portals.",
    "Centralized databases represent single points of failure vulnerable to cyberattacks.",
    "Manual background checks take between 5 to 20 business days to complete.",
    "High administrative costs for universities and candidate screening agencies."
]
for pt in points2:
    add_text_line(tf, "- " + pt, font_size=11.5, color=C_TEXT_MAIN, space_after=10)

# Card 3
add_card(slide2, Inches(8.8), card_y, col_w, card_h, bg_color=C_BLUE_LIGHT, border_color=RGBColor(191, 219, 254))
tb = slide2.shapes.add_textbox(Inches(9.05), card_y + Inches(0.25), col_w - Inches(0.5), card_h - Inches(0.5))
tf = tb.text_frame
tf.word_wrap = True
add_text_line(tf, "The Web3 Opportunity", font_size=16, bold=True, color=C_BLUE, space_after=12)
points3 = [
    "Blockchain delivers an immutable, decentralized ledger immune to unilateral tampering.",
    "Smart contracts enforce deterministic verification rules without human bias.",
    "Public zero-trust verification: Anyone can authenticate credentials in milliseconds.",
    "Cryptographic hashing protects student personal privacy while proving authenticity."
]
for pt in points3:
    add_text_line(tf, "[+] " + pt, font_size=11.5, color=C_NAVY_DARK, space_after=10)


# ==============================================================================
# SLIDE 3: Problem Statement
# ==============================================================================
slide3 = prs.slides.add_slide(blank_layout)
add_header(slide3, "Problem Statement: Current Verification Deficiencies")

quadrants = [
    ("1. Ease of Digital Alteration", "PDF credentials and scans can be altered using basic consumer graphic software. Modifying a student's name, degree class, or course title leaves no digital audit trail.", Inches(0.8), Inches(1.6)),
    ("2. Verification Delay & High Overhead", "Recruiters and institutions must submit email or phone verification requests to universities. This creates substantial delays during critical hiring and admissions cycles.", Inches(6.8), Inches(1.6)),
    ("3. Database Vulnerability & Tampering", "Centralized education repositories are prime targets for ransomware, insider tampering, and server outages. If the server goes down, verification ceases entirely.", Inches(0.8), Inches(4.3)),
    ("4. The Revocation Dilemma", "When a student certificate is cancelled due to academic malpractice, paper copies or cached PDF files remain in circulation with no real-time way for verifiers to discover the revocation.", Inches(6.8), Inches(4.3)),
]

for title, desc, left, top in quadrants:
    add_card(slide3, left, top, Inches(5.6), Inches(2.4))
    tb = slide3.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), Inches(5.1), Inches(2.0))
    tf = tb.text_frame
    tf.word_wrap = True
    add_text_line(tf, title, font_size=14, bold=True, color=C_BLUE, space_after=8)
    add_text_line(tf, desc, font_size=11, color=C_TEXT_MAIN)


# ==============================================================================
# SLIDE 4: Proposed Solution: SkillChain
# ==============================================================================
slide4 = prs.slides.add_slide(blank_layout)
add_header(slide4, "Proposed Solution: SkillChain Architecture")

add_card(slide4, Inches(0.8), Inches(1.6), Inches(4.8), Inches(5.2), bg_color=C_NAVY_DARK, border_color=C_NAVY_DARK)
tb = slide4.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(4.2), Inches(4.6))
tf = tb.text_frame
tf.word_wrap = True
add_text_line(tf, "The SkillChain Philosophy", font_size=18, bold=True, color=C_WHITE, space_after=14)
add_text_line(tf, '"Your Skills. Your Proof."', font_size=15, bold=True, color=RGBColor(96, 165, 250), space_after=14)
add_text_line(tf, "SkillChain is a decentralized skill passport where institutions anchor tamper-proof cryptographic proofs (SHA-256 digests) into Ethereum smart contracts.", font_size=12, color=RGBColor(226, 232, 240), space_after=10)
add_text_line(tf, "Sensitive student records remain off-chain for privacy, while the blockchain guarantees zero-trust immutability and instant public verification.", font_size=12, color=RGBColor(226, 232, 240))

pillars = [
    ("Cryptographic Fingerprinting (SHA-256)", "Generates a deterministic 256-bit hash of student details. Even a single character change completely alters the hash, instantly exposing tampering.", Inches(1.6)),
    ("Immutable On-Chain Ledger", "Certificate records, timestamps, and issuer wallet addresses are recorded in a Solidity contract on Ethereum—immune to centralized manipulation or deletion.", Inches(3.4)),
    ("Instant Public Verification", "Anyone (employers, universities) can query the contract in milliseconds using the Certificate ID without requiring privileged portal logins.", Inches(5.2)),
]

for title, desc, top in pillars:
    add_card(slide4, Inches(5.8), top, Inches(6.8), Inches(1.55))
    tb = slide4.shapes.add_textbox(Inches(6.1), top + Inches(0.18), Inches(6.2), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    add_text_line(tf, title, font_size=13, bold=True, color=C_BLUE, space_after=4)
    add_text_line(tf, desc, font_size=10.5, color=C_TEXT_MAIN)


# ==============================================================================
# SLIDE 5: System Architecture & Workflow
# ==============================================================================
slide5 = prs.slides.add_slide(blank_layout)
add_header(slide5, "System Architecture & End-to-End Workflow")

steps = [
    ("Step 1: Certificate Creation", "Issuer fills out student name, skill title, issuing organization, and date in the React frontend.", Inches(0.8)),
    ("Step 2: SHA-256 & ID Generation", "Frontend normalizes data, generates unique ID (SKC-2026-XXXX), and computes SHA-256 hash using Ethers.js.", Inches(3.8)),
    ("Step 3: Smart Contract Anchoring", "Transaction is signed and broadcasted to SkillChain.sol. The contract records ID, hash, issuer, and timestamp.", Inches(6.8)),
    ("Step 4: Public Verification", "Verifier enters the Certificate ID. The smart contract performs an on-chain hash comparison and flags VERIFIED or INVALID.", Inches(9.8)),
]

for title, desc, left in steps:
    add_card(slide5, left, Inches(1.6), Inches(2.7), Inches(5.2))
    tb = slide5.shapes.add_textbox(left + Inches(0.2), Inches(1.8), Inches(2.3), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    add_text_line(tf, title, font_size=13, bold=True, color=C_BLUE, space_after=10)
    add_text_line(tf, desc, font_size=10.5, color=C_TEXT_MAIN)


# ==============================================================================
# SLIDE 6: Smart Contract Design (SkillChain.sol)
# ==============================================================================
slide6 = prs.slides.add_slide(blank_layout)
add_header(slide6, "Smart Contract Architecture (SkillChain.sol)")

add_card(slide6, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
tb = slide6.shapes.add_textbox(Inches(1.05), Inches(1.8), Inches(5.1), Inches(4.7))
tf = tb.text_frame
tf.word_wrap = True
add_text_line(tf, "Data Structure & Storage", font_size=15, bold=True, color=C_NAVY_DARK, space_after=8)
struct_lines = [
    "struct Certificate {",
    "    string certificateId;",
    "    string certificateHash;  // SHA-256",
    "    address issuer;          // Signer",
    "    uint256 issueTimestamp;  // Block time",
    "    bool isRevoked;",
    "    bool exists;",
    "}",
    "",
    "mapping(string => Certificate) certificates;",
    "event CertificateRegistered(id, hash, issuer);",
    "event CertificateRevoked(id, issuer, time);",
]
for line in struct_lines:
    add_text_line(tf, line, font_size=9.5, color=C_NAVY_LIGHT, space_after=2, font_name="Courier New")

add_card(slide6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
tb = slide6.shapes.add_textbox(Inches(7.05), Inches(1.8), Inches(5.2), Inches(4.7))
tf = tb.text_frame
tf.word_wrap = True
add_text_line(tf, "Core Contract Functions", font_size=15, bold=True, color=C_NAVY_DARK, space_after=10)
methods = [
    ("registerCertificate(id, hash)", "Validates non-empty parameters, rejects duplicate IDs, stores sender as issuer, records block timestamp."),
    ("verifyCertificate(id, hash)", "Compares candidate hash with stored hash (keccak256 comparison). Returns (isValid, isRevoked, exists)."),
    ("revokeCertificate(id)", "Enforces strict authorization: require(cert.issuer == msg.sender). Sets isRevoked = true and emits event."),
    ("getCertificate(id)", "Public view returning full certificate metadata, timestamp, issuer address, and status."),
]
for m_title, m_desc in methods:
    add_text_line(tf, "- " + m_title, font_size=11, bold=True, color=C_BLUE, space_after=2)
    add_text_line(tf, "   " + m_desc, font_size=9.5, color=C_TEXT_MAIN, space_after=6)


# ==============================================================================
# SLIDE 7: Key Project Modules
# ==============================================================================
slide7 = prs.slides.add_slide(blank_layout)
add_header(slide7, "Key Application Modules")

modules = [
    ("Home Landing Page", "Clean, modern entry point reinforcing the project ethos ('Your Skills. Your Proof.'). Provides direct, responsive navigation across all application features.", Inches(0.8), Inches(1.6)),
    ("Add Certificate Module", "Input validation, unique alphanumeric ID generation, SHA-256 payload calculation, and smart contract broadcasting with live transaction hash display.", Inches(6.8), Inches(1.6)),
    ("My Certificates Dashboard", "Overview of credentials issued in the current browser with live smart contract status querying and issuer-only revocation capability.", Inches(0.8), Inches(4.3)),
    ("Public Verification Portal", "Independent verification page. Compares calculated hash with on-chain records; highlights VERIFIED, INVALID, or REVOKED credentials.", Inches(6.8), Inches(4.3)),
]

for title, desc, left, top in modules:
    add_card(slide7, left, top, Inches(5.6), Inches(2.4))
    tb = slide7.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), Inches(5.1), Inches(2.0))
    tf = tb.text_frame
    tf.word_wrap = True
    add_text_line(tf, title, font_size=14, bold=True, color=C_BLUE, space_after=8)
    add_text_line(tf, desc, font_size=11, color=C_TEXT_MAIN)


# ==============================================================================
# SLIDE 8: Security & Cryptographic Mechanisms
# ==============================================================================
slide8 = prs.slides.add_slide(blank_layout)
add_header(slide8, "Security, Privacy & Cryptographic Integrity")

sec_items = [
    ("SHA-256 Data Privacy (Off-Chain Storage)", "Sensitive student personal information (full name, student ID, grades) is hashed prior to blockchain registration. No private student data is stored in public contract state, ensuring compliance with privacy standards while retaining 100% mathematical verifiability.", Inches(1.6)),
    ("Blockchain Immutability & Anti-Tampering", "Once a transaction is confirmed in a block, the certificate record cannot be edited, overwritten, backdated, or deleted by any administrative party or attacker. This guarantees zero-trust tamper resistance.", Inches(3.4)),
    ("Smart Contract Access Control & Duplicate Prevention", "Revocation functions enforce require(cert.issuer == msg.sender), ensuring only the genuine certifying wallet can invalidate credentials. Duplicate registration is eliminated by checking require(!certificates[id].exists).", Inches(5.2)),
]

for title, desc, top in sec_items:
    add_card(slide8, Inches(0.8), top, Inches(11.7), Inches(1.5))
    tb = slide8.shapes.add_textbox(Inches(1.1), top + Inches(0.18), Inches(11.1), Inches(1.14))
    tf = tb.text_frame
    tf.word_wrap = True
    add_text_line(tf, title, font_size=13, bold=True, color=C_BLUE, space_after=4)
    add_text_line(tf, desc, font_size=10.5, color=C_TEXT_MAIN)


# ==============================================================================
# SLIDE 9: Technology Stack
# ==============================================================================
slide9 = prs.slides.add_slide(blank_layout)
add_header(slide9, "Technology Stack & Development Environment")

tech_boxes = [
    ("Frontend Layer", ["React.js 18 (Vite 5)", "JavaScript (ES6+)", "Lucide React (Icons)", "Custom CSS3 (Light/Navy Theme)"], Inches(0.8)),
    ("Smart Contract", ["Solidity ^0.8.20", "EVM Compatible", "Optimized Compiler (Paris)", "Custom Structs & Events"], Inches(3.8)),
    ("Blockchain Network", ["Hardhat Local Network", "Chain ID: 31337", "JSON-RPC Protocol (8545)", "MetaMask Wallet Integration"], Inches(6.8)),
    ("Web3 & Testing", ["Ethers.js (v6)", "Mocha Test Framework", "Chai Assertion Library", "Browser localStorage"], Inches(9.8)),
]

for title, items, left in tech_boxes:
    add_card(slide9, left, Inches(1.6), Inches(2.7), Inches(5.2))
    tb = slide9.shapes.add_textbox(left + Inches(0.2), Inches(1.8), Inches(2.3), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    add_text_line(tf, title, font_size=14, bold=True, color=C_BLUE, space_after=14)
    for item in items:
        add_text_line(tf, "- " + item, font_size=11, color=C_TEXT_MAIN, space_after=8)


# ==============================================================================
# SLIDE 10: Testing & Verification Results
# ==============================================================================
slide10 = prs.slides.add_slide(blank_layout)
add_header(slide10, "Automated Testing & Validation (10/10 Passing)")

add_card(slide10, Inches(0.8), Inches(1.6), Inches(4.2), Inches(5.2), bg_color=RGBColor(240, 253, 244), border_color=RGBColor(187, 247, 208))
tb = slide10.shapes.add_textbox(Inches(1.05), Inches(1.9), Inches(3.7), Inches(4.6))
tf = tb.text_frame
tf.word_wrap = True
add_text_line(tf, "100% Pass Rate", font_size=24, bold=True, color=C_GREEN, space_after=12)
add_text_line(tf, "Test Suite Overview:", font_size=13, bold=True, color=C_NAVY_DARK, space_after=8)
overview_items = [
    "- Framework: Hardhat + Mocha + Chai",
    "- Execution Speed: < 800ms",
    "- Zero Flaky Tests",
    "- Comprehensive Revert Handling",
    "- Real Local Blockchain Simulation"
]
for item in overview_items:
    add_text_line(tf, item, font_size=11, color=C_TEXT_MAIN, space_after=4)

add_card(slide10, Inches(5.3), Inches(1.6), Inches(7.2), Inches(5.2))
tb = slide10.shapes.add_textbox(Inches(5.55), Inches(1.8), Inches(6.7), Inches(4.7))
tf = tb.text_frame
tf.word_wrap = True
add_text_line(tf, "Automated Test Scenarios Covered:", font_size=14, bold=True, color=C_NAVY_DARK, space_after=10)
tests = [
    "[PASS] Registering a certificate on-chain with unique ID and SHA-256 hash",
    "[PASS] Rejecting duplicate certificate IDs (preventing record collisions)",
    "[PASS] Rejecting registration with empty certificate ID or hash parameters",
    "[PASS] Verifying authentic certificate hashes (returns isValid = true)",
    "[PASS] Detecting & rejecting tampered/altered certificate hashes",
    "[PASS] Graceful handling of non-existent certificate IDs",
    "[PASS] Authorizing certificate revocation by genuine issuer address",
    "[PASS] Rejecting revocation attempts by unauthorized non-issuer accounts",
    "[PASS] Rejecting verification for revoked certificates (isRevoked = true)",
    "[PASS] Preventing duplicate revocation attempts on revoked certificates",
]
for t in tests:
    add_text_line(tf, t, font_size=9.8, color=C_TEXT_MAIN, space_after=4)


# ==============================================================================
# SLIDE 11: Live Demonstration Walkthrough
# ==============================================================================
slide11 = prs.slides.add_slide(blank_layout)
add_header(slide11, "Demonstration Walkthrough")

demo_steps = [
    ("1. Issue Certificate", "Navigate to 'Add Certificate', input student name and skill, and click register. Transaction is mined in Hardhat, returning a unique Certificate ID and real Ethereum Tx hash.", Inches(0.8), Inches(1.6)),
    ("2. View in Passport", "Navigate to 'My Certificates'. The issued certificate appears instantly with its issuer, issue date, and an on-chain status verification badge.", Inches(6.8), Inches(1.6)),
    ("3. Public Verification", "Enter the Certificate ID on the public 'Verify' page. The application calculates the SHA-256 hash and compares it on-chain to display a bold VERIFIED badge.", Inches(0.8), Inches(4.3)),
    ("4. Negative & Revoke Test", "Altering even a single character in the certificate details immediately triggers an INVALID alert. Clicking Revoke updates on-chain state to REVOKED.", Inches(6.8), Inches(4.3)),
]

for title, desc, left, top in demo_steps:
    add_card(slide11, left, top, Inches(5.6), Inches(2.4))
    tb = slide11.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), Inches(5.1), Inches(2.0))
    tf = tb.text_frame
    tf.word_wrap = True
    add_text_line(tf, title, font_size=14, bold=True, color=C_BLUE, space_after=8)
    add_text_line(tf, desc, font_size=11, color=C_TEXT_MAIN)


# ==============================================================================
# SLIDE 12: Thank You / Q&A Slide (Dark Theme)
# ==============================================================================
slide12 = prs.slides.add_slide(blank_layout)
bg12 = slide12.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg12.fill.solid()
bg12.fill.fore_color.rgb = C_NAVY_DARK
bg12.line.color.rgb = C_NAVY_DARK
bg12.line.width = Pt(0)

tb12 = slide12.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(10.3), Inches(3.8))
tf12 = tb12.text_frame
tf12.word_wrap = True

add_text_line(tf12, "Thank You!", font_size=44, bold=True, color=C_WHITE, space_after=12)
add_text_line(tf12, "Questions & Answers (Q&A)", font_size=22, bold=True, color=RGBColor(96, 165, 250), space_after=24)
add_text_line(tf12, "GitHub Repository: https://github.com/saksham1126/SkillChain", font_size=13, color=RGBColor(203, 213, 225), space_after=4)
add_text_line(tf12, "Author: Saksham Pathak  |  Project: SkillChain — Blockchain-Based Skill Passport", font_size=13, color=RGBColor(203, 213, 225))

output_filename = "SkillChain_Presentation.pptx"
prs.save(output_filename)
print(f"Clean, uncorrupted presentation saved successfully as: {output_filename}")
